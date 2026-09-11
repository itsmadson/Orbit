"""Letter numbering and rendering.

Two jobs live here:

1. **Numbering** — minting a registered number from a company-defined formula,
   with a counter that resets on the company's own calendar.
2. **Rendering** — turning a letter plus its letterhead into a printable A4
   sheet, as PDF (final) or DOCX (editable).

Both are Persian-first: the calendar defaults to Jalali, text runs right to
left, and numerals can be rendered as Persian digits.
"""

from __future__ import annotations

import io
import re
from datetime import UTC, date, datetime
from html import escape
from pathlib import Path

import jdatetime
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.identity import Company, User
from app.models.office import Letter, Letterhead, LetterNumbering, LetterSequence

FONT_DIR = Path(__file__).resolve().parent.parent / "assets" / "fonts"

FA_DIGITS = "۰۱۲۳۴۵۶۷۸۹"
EN_DIGITS = "0123456789"

JALALI_MONTHS = [
    "فروردین", "اردیبهشت", "خرداد", "تیر", "مرداد", "شهریور",
    "مهر", "آبان", "آذر", "دی", "بهمن", "اسفند",
]

#: Tokens a numbering pattern may use, shown to the user in settings.
NUMBER_TOKENS = {
    "{prefix}": "Company prefix from the numbering settings",
    "{year}": "Four-digit year in the selected calendar (۱۴۰۴)",
    "{yy}": "Two-digit year (۰۴)",
    "{month}": "Two-digit month",
    "{day}": "Two-digit day",
    "{kind}": "Letter-kind code (ص / و / د)",
    "{dept}": "Department code of the author",
    "{seq}": "Counter; pad it as {seq:04} for 0012",
}

_TOKEN_RE = re.compile(r"\{(prefix|year|yy|month|day|kind|dept|seq)(?::0(\d)\|?)?\}")


# --------------------------------------------------------------------------- dates
def to_fa_digits(value: str) -> str:
    return value.translate(str.maketrans(EN_DIGITS, FA_DIGITS))


def to_en_digits(value: str) -> str:
    return value.translate(str.maketrans(FA_DIGITS, EN_DIGITS))


def to_jalali(value: date) -> jdatetime.date:
    return jdatetime.date.fromgregorian(date=value)


def format_letter_date(value: date, calendar: str = "jalali", digits: str = "fa") -> str:
    """`۱۲ مهر ۱۴۰۴` in Jalali, `12 October 2025` in Gregorian."""
    if calendar == "jalali":
        jd = to_jalali(value)
        text = f"{jd.day} {JALALI_MONTHS[jd.month - 1]} {jd.year}"
    else:
        text = value.strftime("%d %B %Y")
    return to_fa_digits(text) if digits == "fa" else text


def numeric_date(value: date, calendar: str = "jalali", digits: str = "fa") -> str:
    """`۱۴۰۴/۰۷/۱۲`."""
    if calendar == "jalali":
        jd = to_jalali(value)
        text = f"{jd.year}/{jd.month:02d}/{jd.day:02d}"
    else:
        text = value.strftime("%Y/%m/%d")
    return to_fa_digits(text) if digits == "fa" else text


# ------------------------------------------------------------------------- numbers
def period_key(numbering: LetterNumbering, when: date) -> str:
    """The bucket a counter belongs to, which is what makes the reset work."""
    if numbering.reset == "never":
        return "-"
    if numbering.calendar == "jalali":
        jd = to_jalali(when)
        year, month = jd.year, jd.month
    else:
        year, month = when.year, when.month
    return f"{year}" if numbering.reset == "yearly" else f"{year}-{month:02d}"


def scope_key(numbering: LetterNumbering, kind: str) -> str:
    return kind if numbering.scope == "per_kind" else "all"


def render_pattern(
    numbering: LetterNumbering,
    *,
    kind: str,
    seq: int,
    when: date,
    dept_code: str = "",
) -> str:
    """Substitute the pattern's tokens. Unknown braces are left untouched."""
    if numbering.calendar == "jalali":
        jd = to_jalali(when)
        year, month, day = jd.year, jd.month, jd.day
    else:
        year, month, day = when.year, when.month, when.day

    values = {
        "prefix": numbering.prefix or "",
        "year": f"{year}",
        "yy": f"{year % 100:02d}",
        "month": f"{month:02d}",
        "day": f"{day:02d}",
        "kind": (numbering.kind_codes or {}).get(kind, ""),
        "dept": dept_code or "",
        "seq": f"{seq}",
    }

    def substitute(match: re.Match[str]) -> str:
        token, pad = match.group(1), match.group(2)
        raw = values.get(token, "")
        if token == "seq" and pad:
            raw = f"{seq:0{int(pad)}d}"
        return raw

    text = _TOKEN_RE.sub(substitute, numbering.pattern or "{seq}")
    # A pattern with an empty token (no prefix, no dept) must not leave "//".
    sep = re.escape(numbering.separator or "/")
    text = re.sub(rf"({sep})+", numbering.separator or "/", text).strip(numbering.separator or "/")
    return to_fa_digits(text) if numbering.digits == "fa" else text


def preview_number(numbering: LetterNumbering, *, kind: str = "outgoing",
                   when: date | None = None, seq: int | None = None) -> str:
    when = when or date.today()
    return render_pattern(numbering, kind=kind, seq=seq if seq is not None else 12, when=when)


def default_numbering(db: Session, company_id) -> LetterNumbering | None:
    return db.scalar(
        select(LetterNumbering)
        .where(LetterNumbering.company_id == company_id)
        .order_by(LetterNumbering.is_default.desc(), LetterNumbering.created_at)
    )


def allocate_number(
    db: Session,
    *,
    numbering: LetterNumbering,
    kind: str,
    when: date,
    dept_code: str = "",
) -> tuple[str, int, str]:
    """Consume the next counter value and render it. Returns (number, seq, period).

    The counter row is locked for the transaction, so two clerks registering a
    letter at the same moment cannot receive the same number.
    """
    scope = scope_key(numbering, kind)
    period = period_key(numbering, when)

    row = db.scalar(
        select(LetterSequence)
        .where(
            LetterSequence.numbering_id == numbering.id,
            LetterSequence.scope_key == scope,
            LetterSequence.period_key == period,
        )
        .with_for_update()
    )
    if row is None:
        row = LetterSequence(
            company_id=numbering.company_id, numbering_id=numbering.id,
            scope_key=scope, period_key=period, next_value=numbering.start_at or 1,
        )
        db.add(row)
        db.flush()

    seq = row.next_value
    row.next_value = seq + 1
    number = render_pattern(numbering, kind=kind, seq=seq, when=when, dept_code=dept_code)
    return number, seq, period


# ----------------------------------------------------------------------- templates
_VAR_RE = re.compile(r"\{\{\s*([a-zA-Z_][a-zA-Z0-9_]*)\s*\}\}")


def fill(text: str | None, variables: dict) -> str:
    """Replace {{name}} placeholders; unknown ones are left visible on purpose,
    so a half-filled template is obvious rather than silently blank."""
    if not text:
        return ""
    return _VAR_RE.sub(lambda m: str(variables.get(m.group(1), m.group(0))), text)


def template_variables(letter: Letter, company: Company, author: User | None) -> dict:
    return {
        "recipient": letter.recipient_name or "",
        "recipient_title": letter.recipient_title or "",
        "recipient_org": letter.recipient_org or "",
        "subject": letter.subject or "",
        "number": letter.number or "",
        "date": numeric_date(letter.letter_date),
        "company": company.name,
        "sender": letter.sender_name or (author.full_name if author else ""),
        "sender_title": letter.sender_title or (author.title if author else ""),
    }


# ------------------------------------------------------------------------ rendering
def _contact_line(head: Letterhead, digits: str) -> str:
    bits = [b for b in [head.address, head.phone, head.email, head.website] if b]
    line = " · ".join(escape(b) for b in bits)
    return to_fa_digits(line) if digits == "fa" else line


def letter_html(
    letter: Letter,
    head: Letterhead | None,
    company: Company,
    *,
    author: User | None = None,
    signer: User | None = None,
    standalone: bool = True,
) -> str:
    """The single source of truth for how a letter looks — used by the PDF
    renderer and by the browser preview, so what you see is what prints."""
    head = head or Letterhead(
        company_id=company.id, name="default", org_name=company.name,
    )
    digits = "fa" if head.language == "fa" else "en"
    rtl = head.direction == "rtl"
    calendar = "jalali" if head.language == "fa" else "gregorian"

    def num(value: str) -> str:
        return to_fa_digits(value) if digits == "fa" else value

    labels = (
        {"number": "شماره", "date": "تاریخ", "attachment": "پیوست",
         "subject": "موضوع", "follow": "پیرو", "page": "صفحه"}
        if rtl else
        {"number": "No.", "date": "Date", "attachment": "Enclosure",
         "subject": "Subject", "follow": "Ref.", "page": "Page"}
    )

    org_block = f"""
      <div class="org">
        {f'<img class="logo" src="{head.logo_data_url}" alt="" />' if head.logo_data_url else ''}
        <div class="org-text">
          <div class="org-name">{escape(head.org_name or company.name)}</div>
          {f'<div class="org-sub">{escape(head.org_subtitle)}</div>' if head.org_subtitle else ''}
        </div>
      </div>"""

    header = head.header_html or org_block

    meta_rows = [
        (labels["number"], num(letter.number or "—")),
        (labels["date"], numeric_date(letter.letter_date, calendar, digits)),
    ]
    if letter.attachment_note:
        meta_rows.append((labels["attachment"], escape(letter.attachment_note)))
    if letter.follow_up_of:
        meta_rows.append((labels["follow"], num(letter.follow_up_of)))

    meta = "".join(
        f'<div class="meta-row"><span class="meta-label">{k}</span>'
        f'<span class="meta-sep">:</span><span class="meta-value">{v}</span></div>'
        for k, v in meta_rows
    )

    marks = []
    if letter.confidentiality != "normal":
        marks.append("محرمانه" if letter.confidentiality == "confidential" else "سری")
    if letter.urgency != "normal":
        marks.append("فوری" if letter.urgency == "urgent" else "آنی")
    marks_html = (
        f'<div class="marks">{" · ".join(escape(m) for m in marks)}</div>' if marks else ""
    )

    recipient_lines = [letter.salutation or ""]
    if letter.recipient_org and letter.recipient_org not in (letter.salutation or ""):
        recipient_lines.append(letter.recipient_org)
    recipient = "".join(
        f'<div class="recipient-line">{escape(line)}</div>' for line in recipient_lines if line
    )

    signature_name = escape(
        letter.sender_name or (signer.full_name if signer else "") or
        (author.full_name if author else "")
    )
    signature_title = escape(
        letter.sender_title or (signer.title if signer else "") or
        (author.title if author else "")
    )
    signature_img = (
        f'<img class="sig-img" src="{head.signature_data_url}" alt="" />'
        if head.signature_data_url and letter.status in ("signed", "sent", "archived") else ""
    )
    stamp_img = (
        f'<img class="stamp" src="{head.stamp_data_url}" alt="" />'
        if head.stamp_data_url and letter.status in ("signed", "sent", "archived") else ""
    )

    cc_html = ""
    if letter.cc:
        items = "".join(f"<li>{escape(str(c))}</li>" for c in letter.cc)
        cc_html = f'<div class="cc"><span>رونوشت:</span><ul>{items}</ul></div>'

    footer = head.footer_html or f'<div class="contact">{_contact_line(head, digits)}</div>'

    body = f"""
    <div class="sheet">
      <header class="letter-header">
        {header}
        <div class="meta">{meta}</div>
      </header>
      {marks_html}
      <section class="to">{recipient}</section>
      <h1 class="subject"><span class="subject-label">{labels['subject']}:</span> {escape(letter.subject)}</h1>
      <article class="body">{letter.body or ''}</article>
      {f'<div class="closing">{escape(letter.closing)}</div>' if letter.closing else ''}
      <div class="signature">
        {signature_img}
        <div class="sig-name">{signature_name}</div>
        <div class="sig-title">{signature_title}</div>
        {stamp_img}
      </div>
      {cc_html}
      <footer class="letter-footer">{footer}</footer>
    </div>"""

    if not standalone:
        return body

    page_rule = f"""
      @page {{
        size: {head.paper or 'A4'};
        margin: {head.margin_top_mm}mm {head.margin_x_mm}mm {head.margin_bottom_mm}mm;
        {'@bottom-center { content: counter(page) " / " counter(pages); font-size: 8pt; color: #888; }'
         if head.show_page_numbers else ''}
      }}"""

    return f"""<!doctype html>
<html lang="{head.language}" dir="{head.direction}">
<head>
<meta charset="utf-8" />
<title>{escape(letter.subject)}</title>
<style>
  @font-face {{
    font-family: "Vazirmatn";
    src: url("file://{FONT_DIR}/Vazirmatn-Regular.ttf") format("truetype");
    font-weight: 400;
  }}
  @font-face {{
    font-family: "Vazirmatn";
    src: url("file://{FONT_DIR}/Vazirmatn-Bold.ttf") format("truetype");
    font-weight: 700;
  }}
  {page_rule}
  * {{ box-sizing: border-box; }}
  body {{
    font-family: "{head.font_family}", "Vazirmatn", sans-serif;
    font-size: {head.font_size_pt}pt;
    line-height: 1.9;
    color: #14181f;
    margin: 0;
    direction: {head.direction};
  }}
  .letter-header {{
    display: flex; align-items: flex-start; justify-content: space-between;
    gap: 16mm; padding-bottom: 4mm; border-bottom: 1.5pt solid {head.accent_color};
  }}
  .org {{ display: flex; align-items: center; gap: 4mm; }}
  .logo {{ height: 16mm; width: auto; }}
  .org-name {{ font-size: {head.font_size_pt + 3}pt; font-weight: 700; }}
  .org-sub {{ font-size: {head.font_size_pt - 2}pt; color: #5b6270; }}
  .meta {{ min-width: 42mm; font-size: {head.font_size_pt - 1}pt; }}
  .meta-row {{ display: flex; gap: 1mm; }}
  .meta-label {{ min-width: 14mm; color: #5b6270; }}
  .meta-value {{ font-weight: 600; }}
  .marks {{
    margin-top: 4mm; display: inline-block; padding: 1mm 3mm;
    border: 1pt solid #c62828; color: #c62828; font-size: {head.font_size_pt - 2}pt;
    font-weight: 700; border-radius: 2mm;
  }}
  .to {{ margin-top: 8mm; font-weight: 600; }}
  .recipient-line {{ line-height: 1.7; }}
  .subject {{
    margin: 6mm 0 4mm; font-size: {head.font_size_pt + 1}pt; font-weight: 700;
  }}
  .subject-label {{ color: #5b6270; font-weight: 500; }}
  .body {{ text-align: justify; }}
  .body p {{ margin: 0 0 3mm; text-indent: 6mm; }}
  .body ul, .body ol {{ padding-inline-start: 8mm; margin: 0 0 3mm; }}
  .body table {{ width: 100%; border-collapse: collapse; margin: 3mm 0; }}
  .body th, .body td {{ border: 0.5pt solid #b9c0cc; padding: 1.5mm 2mm; }}
  .closing {{ margin-top: 5mm; }}
  .signature {{
    margin-top: 10mm; text-align: center; width: 60mm;
    margin-inline-start: auto; position: relative;
  }}
  .sig-img {{ height: 16mm; display: block; margin: 0 auto 1mm; }}
  .sig-name {{ font-weight: 700; }}
  .sig-title {{ font-size: {head.font_size_pt - 2}pt; color: #5b6270; }}
  .stamp {{ height: 24mm; position: absolute; inset-inline-start: -18mm; top: -4mm; opacity: 0.85; }}
  .cc {{ margin-top: 10mm; font-size: {head.font_size_pt - 2}pt; color: #5b6270; }}
  .cc ul {{ margin: 1mm 0 0; padding-inline-start: 6mm; }}
  .letter-footer {{
    position: fixed; bottom: 0; inset-inline: 0;
    border-top: 0.5pt solid #d6dae2; padding-top: 2mm;
    font-size: {head.font_size_pt - 3}pt; color: #6b7280; text-align: center;
  }}
</style>
</head>
<body>{body}</body>
</html>"""


def render_pdf(html: str) -> bytes:
    from weasyprint import HTML  # imported lazily: heavy native dependency

    return HTML(string=html, base_url=str(FONT_DIR)).write_pdf()


def render_docx(
    letter: Letter,
    head: Letterhead | None,
    company: Company,
    *,
    author: User | None = None,
    signer: User | None = None,
) -> bytes:
    """An editable copy. Shapes the document RTL so Word opens it correctly."""
    from docx import Document as Docx
    from docx.enum.section import WD_SECTION
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt, RGBColor

    head = head or Letterhead(company_id=company.id, name="default", org_name=company.name)
    rtl = head.direction == "rtl"
    digits = "fa" if head.language == "fa" else "en"
    calendar = "jalali" if head.language == "fa" else "gregorian"

    doc = Docx()
    section = doc.sections[0]
    section.top_margin = Mm(head.margin_top_mm)
    section.bottom_margin = Mm(head.margin_bottom_mm)
    section.left_margin = section.right_margin = Mm(head.margin_x_mm)
    if rtl:
        bidi = section._sectPr.makeelement(qn("w:bidi"), {})
        section._sectPr.append(bidi)

    style = doc.styles["Normal"]
    style.font.name = head.font_family or "Vazirmatn"
    style.font.size = Pt(head.font_size_pt)
    style.element.rPr.rFonts.set(qn("w:cs"), head.font_family or "Vazirmatn")

    def para(text="", *, bold=False, size=None, align=None, space_after=6):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(space_after)
        if rtl:
            p.paragraph_format.alignment = align or WD_ALIGN_PARAGRAPH.RIGHT
            pPr = p._p.get_or_add_pPr()
            pPr.append(pPr.makeelement(qn("w:bidi"), {}))
        else:
            p.paragraph_format.alignment = align or WD_ALIGN_PARAGRAPH.LEFT
        if text:
            run = p.add_run(text)
            run.bold = bold
            run.font.size = Pt(size or head.font_size_pt)
            run.font.name = head.font_family or "Vazirmatn"
            run._element.rPr.rFonts.set(qn("w:cs"), head.font_family or "Vazirmatn")
            if rtl:
                run._element.rPr.append(run._element.rPr.makeelement(qn("w:rtl"), {}))
        return p

    def num(value: str) -> str:
        return to_fa_digits(value) if digits == "fa" else value

    para(head.org_name or company.name, bold=True, size=head.font_size_pt + 3,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    if head.org_subtitle:
        para(head.org_subtitle, size=head.font_size_pt - 2,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)

    label_number = "شماره" if rtl else "No."
    label_date = "تاریخ" if rtl else "Date"
    label_attach = "پیوست" if rtl else "Enclosure"
    meta = f"{label_number}: {num(letter.number or '—')}    " \
           f"{label_date}: {numeric_date(letter.letter_date, calendar, digits)}"
    if letter.attachment_note:
        meta += f"    {label_attach}: {letter.attachment_note}"
    para(meta, size=head.font_size_pt - 1, space_after=12)

    if letter.confidentiality != "normal" or letter.urgency != "normal":
        marks = []
        if letter.confidentiality != "normal":
            marks.append("محرمانه" if letter.confidentiality == "confidential" else "سری")
        if letter.urgency != "normal":
            marks.append("فوری" if letter.urgency == "urgent" else "آنی")
        mark = para(" · ".join(marks), bold=True, size=head.font_size_pt - 1)
        mark.runs[0].font.color.rgb = RGBColor(0xC6, 0x28, 0x28)

    if letter.salutation:
        para(letter.salutation, bold=True, space_after=2)
    if letter.recipient_org and letter.recipient_org not in (letter.salutation or ""):
        para(letter.recipient_org, space_after=10)

    subject_label = "موضوع" if rtl else "Subject"
    para(f"{subject_label}: {letter.subject}", bold=True, space_after=10)

    for block in html_to_blocks(letter.body or ""):
        para(block, space_after=6)

    if letter.closing:
        para(letter.closing, space_after=14)

    signer_name = letter.sender_name or (signer.full_name if signer else "") or \
        (author.full_name if author else "")
    signer_title = letter.sender_title or (signer.title if signer else "") or \
        (author.title if author else "")
    para(signer_name, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
    if signer_title:
        para(signer_title, size=head.font_size_pt - 2,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)

    if letter.cc:
        para(("رونوشت: " if rtl else "CC: ") + "، ".join(str(c) for c in letter.cc),
             size=head.font_size_pt - 2)

    contact = " · ".join(b for b in [head.address, head.phone, head.email, head.website] if b)
    if contact:
        para(contact, size=head.font_size_pt - 3, align=WD_ALIGN_PARAGRAPH.CENTER)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


_BLOCK_RE = re.compile(r"<(?:p|div|li|h[1-6])[^>]*>(.*?)</(?:p|div|li|h[1-6])>", re.S | re.I)
_TAG_RE = re.compile(r"<[^>]+>")


def html_to_blocks(html: str) -> list[str]:
    """Flatten the editor's HTML into paragraphs for the DOCX writer."""
    if not html:
        return []
    html = re.sub(r"<br\s*/?>", "\n", html, flags=re.I)
    blocks = _BLOCK_RE.findall(html)
    if not blocks:
        blocks = [html]
    out: list[str] = []
    for block in blocks:
        text = _TAG_RE.sub("", block)
        text = (text.replace("&nbsp;", " ").replace("&amp;", "&")
                .replace("&lt;", "<").replace("&gt;", ">").replace("&quot;", '"'))
        text = text.strip()
        if text:
            out.append(text)
    return out


def download_filename(letter: Letter, extension: str) -> str:
    """A Content-Disposition value carrying both an ASCII fallback and the real
    Persian filename, because HTTP headers are latin-1 but the number is not."""
    from urllib.parse import quote

    number = letter.number or "letter"
    ascii_name = re.sub(r"[^A-Za-z0-9._-]+", "-", to_en_digits(number)).strip("-")
    if not ascii_name or ascii_name in {"-", "."}:
        ascii_name = f"letter-{letter.number_seq or ''}".strip("-") or "letter"
    utf8_name = f"{number}.{extension}".replace("/", "-")
    return (
        f'attachment; filename="{ascii_name}.{extension}"; '
        f"filename*=UTF-8''{quote(utf8_name)}"
    )


def artifact_key(letter: Letter, extension: str) -> str:
    safe = re.sub(r"[^A-Za-z0-9_-]+", "-", to_en_digits(letter.number or str(letter.id)))
    return f"letters/{letter.company_id}/{letter.id}/{safe or 'letter'}.{extension}"


def mark_rendered(letter: Letter) -> None:
    letter.rendered_at = datetime.now(UTC)
